"""The ingestion report as an Excel workbook, for filtering and sorting.

The same sources as the Markdown report (load result, profile, gate, claims,
the masked preview), so the two can never disagree. Values in personal-data
and identifier columns appear only as their shape, as in every other output.
The rejected lines are the one exception, shown verbatim for the same reason
as in `rejected_lines.csv`: the point is to find them in the source file.

Every string is written as text, never as a formula. A rejected line that
starts with `=` must not run when the workbook is opened.
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pipeline.contracts.audit import GateResult
from pipeline.contracts.context_brief import ClaimResult
from pipeline.contracts.ingestion import LoadResult
from pipeline.contracts.profile import DatasetProfile
from pipeline.ingestion.compact import compact_column
from pipeline.ingestion.pii import mask_shape
from pipeline.report import Preview, masked_columns

XLSX_FILENAME = "ingestion_report.xlsx"
SHEET_NAMES = (
    "Summary",
    "Gate 1 checks",
    "Flags",
    "Columns",
    "Claims",
    "Rejected lines",
    "Renames",
    "Preview",
)
MAX_COLUMN_WIDTH = 60

_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
_OUTCOME_FILL = {
    "pass": PatternFill("solid", fgColor="C6EFCE"),
    "flag": PatternFill("solid", fgColor="FFEB9C"),
    "block": PatternFill("solid", fgColor="FFC7CE"),
}


def write_ingestion_xlsx(
    path: str | Path,
    load_result: LoadResult,
    profile: DatasetProfile,
    gate: GateResult,
    claim_results: Sequence[ClaimResult],
    preview: Preview,
) -> Path:
    """Write the whole report as one workbook. Every sheet exists on every run."""
    book = Workbook()
    book.remove(book.active)
    sheets = {name: book.create_sheet(name) for name in SHEET_NAMES}

    _summary(sheets["Summary"], load_result, profile, gate)
    _gate_checks(sheets["Gate 1 checks"], gate)
    _flags(sheets["Flags"], gate)
    _columns(sheets["Columns"], profile)
    _claims(sheets["Claims"], claim_results)
    _rejects(sheets["Rejected lines"], load_result)
    _renames(sheets["Renames"], load_result)
    _preview(sheets["Preview"], preview)

    for sheet in sheets.values():
        _as_text(sheet)
        _fit_widths(sheet)

    target = Path(path)
    book.save(target)
    return target


def _summary(
    sheet: Worksheet, load_result: LoadResult, profile: DatasetProfile, gate: GateResult
) -> None:
    file = load_result.file
    total = load_result.n_source_rows
    rows: list[tuple[str, object]] = [
        ("Gate 1", gate.outcome.upper()),
        ("File", Path(file.path).name),
        ("Format", file.format),
        ("Encoding", file.encoding or ""),
        ("Delimiter", file.delimiter or ""),
        ("Sheet", file.sheet or ""),
        ("Size (bytes)", file.size_bytes),
        ("File SHA-256", file.file_hash),
        ("Data SHA-256", load_result.data_hash),
        ("Data rows in file", total),
        ("Loaded", load_result.n_loaded),
        ("Rejected", load_result.n_rejected),
        ("Reject rate", load_result.n_rejected / total if total else 0.0),
        ("Columns", len(load_result.columns)),
        ("Duplicate rows", profile.n_duplicate_rows),
        ("Candidate keys", ", ".join(profile.candidate_keys)),
        ("Empty columns", ", ".join(load_result.empty_columns)),
        ("Profiled on a sample", "yes" if profile.sampled else "no"),
        ("Loaded at (UTC)", load_result.loaded_at.strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for label, value in rows:
        sheet.append([(label), (value)])
        sheet.cell(sheet.max_row, 1).font = _BOLD
    sheet["B1"].fill = _OUTCOME_FILL[gate.outcome]
    sheet["B1"].font = _BOLD
    reject_rate_row = next(i for i, (label, _) in enumerate(rows, 1) if label == "Reject rate")
    sheet.cell(reject_rate_row, 2).number_format = "0.00%"

    if file.detection_notes:
        sheet.append([])
        sheet.append([("How the file was interpreted")])
        sheet.cell(sheet.max_row, 1).font = _BOLD
        for note in file.detection_notes:
            sheet.append(["", (note)])


def _gate_checks(sheet: Worksheet, gate: GateResult) -> None:
    _header(sheet, ["Check", "Result", "Severity", "Message"])
    for check in gate.checks:
        result = "pass" if check.passed else "FAIL"
        sheet.append([(check.name), result, check.severity, (check.message)])
        if not check.passed:
            sheet.cell(sheet.max_row, 2).fill = _OUTCOME_FILL["block"]


def _flags(sheet: Worksheet, gate: GateResult) -> None:
    _header(sheet, ["Severity", "Code", "Column", "Message"])
    order = {"error": 0, "warn": 1, "info": 2}
    for flag in sorted(gate.flags, key=lambda f: order.get(f.severity, 3)):
        sheet.append([flag.severity, (flag.code), (flag.column or ""), (flag.message)])


def _columns(sheet: Worksheet, profile: DatasetProfile) -> None:
    _header(
        sheet,
        [
            "Column",
            "Position",
            "Read as",
            "Ambiguous",
            "Fits",
            "Missing",
            "Missing share",
            "Blank",
            "Distinct",
            "Distinct share",
            "Min",
            "Median",
            "Max",
            "Mean",
            "Far outliers",
            "Length min",
            "Length max",
            "Null tokens",
            "Possible coded missing",
            "Personal data",
            "Masked",
            "Example values",
            "Did not fit",
            "Format",
            "Notes",
        ],
    )
    masked = masked_columns(profile)
    for column in profile.columns:
        entry = compact_column(column)
        hide = column.name in masked
        numeric = None if hide else column.numeric
        sentinels = [s.value for s in column.sentinel_candidates]
        pii = entry.get("pii")
        sheet.append(
            [
                (column.name),
                column.position,
                entry["type"],
                "yes" if entry.get("ambiguous") else "",
                column.type_proposal.coverage,
                column.n_missing,
                column.missing_share,
                column.n_blank,
                column.n_distinct,
                column.distinct_share,
                numeric.min if numeric else None,
                numeric.median if numeric else None,
                numeric.max if numeric else None,
                numeric.mean if numeric else None,
                numeric.n_outliers if numeric else None,
                column.length.min if column.length else None,
                column.length.max if column.length else None,
                (_join(f"{k or '(blank)'}: {v}" for k, v in column.null_tokens.items())),
                (_join(mask_shape(s) if hide else s for s in sentinels)),
                (f"{pii['level']}: {', '.join(pii['kinds'])}" if pii else ""),
                "yes" if hide else "",
                (_join(entry.get("examples", []))),
                (_join(entry.get("unparsed_examples", []))),
                (_join(f"{k}={v}" for k, v in entry.get("format", {}).items())),
                (_join(entry.get("notes", []), sep="; ")),
            ]
        )
        row = sheet.max_row
        for col in (5, 7, 10):
            sheet.cell(row, col).number_format = "0.0%"
    sheet.freeze_panes = "B2"


def _claims(sheet: Worksheet, claim_results: Sequence[ClaimResult]) -> None:
    _header(sheet, ["Claim", "Kind", "Column", "Result", "Message"])
    label = {"pass": "holds", "fail": "does not hold", "not_checkable": "could not be checked"}
    for result in claim_results:
        claim = result.claim
        sheet.append(
            [
                (claim.text or claim.kind),
                (claim.kind),
                (getattr(claim, "column", None) or ""),
                label[result.status],
                (result.message),
            ]
        )
        if result.status == "fail":
            sheet.cell(sheet.max_row, 4).fill = _OUTCOME_FILL["block"]


def _rejects(sheet: Worksheet, load_result: LoadResult) -> None:
    _header(sheet, ["Line in file", "Reason", "Text (unmasked, may contain personal data)"])
    for reject in load_result.rejects:
        sheet.append([reject.source_line, (reject.reason), (reject.raw)])


def _renames(sheet: Worksheet, load_result: LoadResult) -> None:
    _header(sheet, ["Position", "In the file", "Now called", "Why"])
    for rename in load_result.renames:
        sheet.append([rename.position, (repr(rename.original)), (rename.final), (rename.reason)])


def _preview(sheet: Worksheet, preview: Preview) -> None:
    _header(sheet, ["Line in file", *preview.columns])
    for source_line, values in preview.rows:
        sheet.append([source_line, *((v) for v in values)])
    if preview.masked:
        sheet.append([])
        sheet.append(
            [("Shown as shape only (personal data or identifier): ") + ", ".join(preview.masked)]
        )


def _header(sheet: Worksheet, names: Sequence[str]) -> None:
    sheet.append([(n) for n in names])
    for cell in sheet[1]:
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(names))}1"


def _fit_widths(sheet: Worksheet) -> None:
    for column in sheet.columns:
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(max(longest + 2, 8), MAX_COLUMN_WIDTH)
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=longest > MAX_COLUMN_WIDTH)


def _join(values, sep: str = ", ") -> str:
    return sep.join(str(v) for v in values)


def _as_text(sheet: Worksheet) -> None:
    """Store every string as text. openpyxl otherwise writes '=...' as a formula."""
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"


__all__ = ["SHEET_NAMES", "XLSX_FILENAME", "write_ingestion_xlsx"]

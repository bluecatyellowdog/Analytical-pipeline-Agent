"""The Excel report: same numbers as the Markdown report, in a form you can filter."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from pipeline.ingest import ingest
from pipeline.ingestion.raw import RawStore
from pipeline.layout import DATA_COLUMNS_XLSX, run_file
from pipeline.report_xlsx import SHEET_NAMES


@pytest.fixture
def run(fixtures_dir, tmp_path):
    def _run(name: str):
        with RawStore.open() as store:
            result = ingest(fixtures_dir / name, store=store)
        out = result.write(tmp_path / name)
        return result, load_workbook(run_file(out, DATA_COLUMNS_XLSX))

    return _run


class TestWorkbook:
    def test_every_run_writes_the_same_sheets(self, run) -> None:
        _, clean = run("danish_semicolon.csv")
        _, broken = run("too_many_fields.csv")
        assert clean.sheetnames == list(SHEET_NAMES)
        assert broken.sheetnames == list(SHEET_NAMES)

    def test_summary_matches_the_load_result(self, run) -> None:
        result, book = run("too_many_fields.csv")
        cells = {row[0].value: row[1].value for row in book["Summary"].iter_rows(min_row=1)}
        assert cells["Gate 1"] == "BLOCK"
        assert cells["Data rows in file"] == result.load_result.n_source_rows
        assert cells["Loaded"] == result.load_result.n_loaded
        assert cells["Rejected"] == result.load_result.n_rejected

    def test_one_row_per_column(self, run) -> None:
        result, book = run("danish_semicolon.csv")
        names = [r[0].value for r in book["Columns"].iter_rows(min_row=2)]
        assert names == result.profile.column_names

    def test_gate_checks_are_listed_with_their_outcome(self, run) -> None:
        result, book = run("too_many_fields.csv")
        rows = {r[0].value: r[1].value for r in book["Gate 1 checks"].iter_rows(min_row=2)}
        assert len(rows) == len(result.gate.checks)
        assert rows["reject_rate_within_threshold"] == "FAIL"
        assert rows["unique_column_names"] == "pass"

    def test_rejected_lines_are_all_there(self, run) -> None:
        result, book = run("too_many_fields.csv")
        rows = [[c.value for c in r] for r in book["Rejected lines"].iter_rows(min_row=2)]
        assert [r[0] for r in rows] == [r.source_line for r in result.load_result.rejects]
        assert all(r[1] and r[2] for r in rows)

    def test_no_rejects_leaves_a_header_only_sheet(self, run) -> None:
        _, book = run("danish_semicolon.csv")
        assert book["Rejected lines"].max_row == 1


class TestSafety:
    def test_personal_values_are_masked_in_the_preview_and_examples(self, run) -> None:
        _, book = run("cpr_numbers.csv")
        text = " ".join(
            str(cell.value)
            for sheet in ("Preview", "Columns")
            for row in book[sheet].iter_rows()
            for cell in row
            if cell.value is not None
        )
        for cpr in ("010190-1234", "311299-5678", "290204-1111", "150385-9999"):
            assert cpr not in text

    def test_formula_like_text_is_stored_as_text(self, tmp_path) -> None:
        path = tmp_path / "evil.csv"
        rows = "\n".join(f"{i},x" for i in range(300))
        path.write_text(f"id,navn\n{rows}\n=1+1,a,b\n", encoding="utf-8")
        with RawStore.open() as store:
            result = ingest(path, store=store)
        out = result.write(tmp_path / "out")
        book = load_workbook(run_file(out, DATA_COLUMNS_XLSX))
        cell = book["Rejected lines"]["C2"]
        assert cell.data_type == "s"
        assert cell.value.startswith("=1+1")

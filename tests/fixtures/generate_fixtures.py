"""Generate the fixture set from spec 4.5.

The files are committed, not generated at test time: tests must run against
exactly the bytes a reviewer can open. This script only regenerates them, and
writes bytes explicitly so encodings and line endings are never left to chance.

Run: uv run python tests/fixtures/generate_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent


def write(name: str, data: bytes) -> None:
    (HERE / name).write_bytes(data)
    print(f"wrote {name} ({len(data)} bytes)")


def lines(*rows: str, newline: str = "\n") -> str:
    return newline.join(rows) + newline


# 1 — the headline Danish export: semicolons, decimal commas, leading-zero IDs,
# NA tokens, a 999 sentinel and an email column.
def danish_semicolon() -> None:
    text = lines(
        "kunde_id;alder;beloeb;dato;email",
        "00417;31;12.345,50;03-04-2021;anna.hansen@example.dk",
        "00418;NA;999;15-07-2021;bo@example.dk",
        "00419;54;7.800,00;28-02-2020;",
        "00420;23;250.000,00;13-11-2022;carla@example.dk",
        "00421;67;999;30-01-2019;",
        "00422;45;1.234,50;17-06-2023;erik@example.dk",
        "00423;NA;56.000,00;21-09-2021;NA",
        "00424;38;3.999,95;14-12-2020;freja@example.dk",
        "00425;29;180.500,00;25-05-2024;",
        "00426;71;999;19-08-2019;gorm@example.dk",
        "00427;19;15.000,00;16-03-2022;hanne@example.dk",
        "00428;52;92.750,25;27-10-2023;",
    )
    write("danish_semicolon.csv", text.encode("utf-8"))


# 2 — a typical Excel "Save as CSV": UTF-8 with a byte-order mark and CRLF.
# The BOM must not end up inside the first column's name.
def bom_excel_export() -> None:
    text = lines(
        "id,navn,antal",
        "1,Alpha,10",
        "2,Beta,20",
        "3,Gamma,30",
        newline="\r\n",
    )
    write("bom_excel_export.csv", b"\xef\xbb\xbf" + text.encode("utf-8"))


# 3 — an older Danish system export: cp1252, with aeoeaa in headers and values.
def latin1_aeoeaa() -> None:
    text = lines(
        "navn;by;beløb",
        "Søren Østergård;København;1.200,00",
        "Åse Bæk;Ærøskøbing;850,50",
        "Mads Nørgaard;Århus;2.300,00",
    )
    write("latin1_aeoeaa.csv", text.encode("cp1252"))


# 4 — a quoted field containing a newline: one logical row spans two lines, and
# the row count must follow the records, not the lines.
def quoted_newline() -> None:
    text = (
        "id,kommentar,beloeb\n"
        '1,"Kunden ringede.\nSagen er lukket.",100\n'
        "2,Ingen bemærkninger,200\n"
        '3,"Bemærk: ""vigtig"" sag",300\n'
    )
    write("quoted_newline.csv", text.encode("utf-8"))


# 5 — duplicate and blank column headers, which must be renamed deterministically
# and reported, never silently collapsed.
def duplicate_blank_headers() -> None:
    text = lines(
        "id,navn,,navn,",
        "1,Alpha,x,Alpha Corp,y",
        "2,Beta,z,Beta Ltd,w",
    )
    write("duplicate_blank_headers.csv", text.encode("utf-8"))


# 6 — malformed lines. Both must be rejected with a reason and a source line,
# and both must remain visible in the rejects table.
def too_many_fields() -> None:
    text = lines(
        "id,navn,beloeb",
        "1,Alpha,100",
        "2,Beta,200,EXTRA",
        "3,Gamma,300",
        "4,Delta",
        "5,Epsilon,500",
    )
    write("too_many_fields.csv", text.encode("utf-8"))


# 7 — every day is 12 or below, so day/month order cannot be determined from the
# data alone. Must be flagged as ambiguous, not guessed.
def ambiguous_dates() -> None:
    text = lines(
        "id,dato",
        "1,03/04/2021",
        "2,11/12/2021",
        "3,01/02/2022",
        "4,07/09/2022",
        "5,12/05/2023",
        "6,02/10/2023",
    )
    write("ambiguous_dates.csv", text.encode("utf-8"))


# 8 — Danish CPR numbers. Only the ones whose first six digits form a real date
# count as CPR; the rest are look-alikes and must not raise the PII level alone.
def cpr_numbers() -> None:
    text = lines(
        "sagsnr;cpr;note",
        "1;010190-1234;gyldig",
        "2;311299-5678;gyldig",
        "3;290204-1111;gyldig skudaar",
        "4;150385-9999;gyldig",
        "5;320199-1234;ugyldig dato",
        "6;001390-4321;ugyldig maaned",
        "7;123456-7890;ikke en dato",
        "8;;mangler",
    )
    write("cpr_numbers.csv", text.encode("utf-8"))


# 9 — a two-sheet workbook whose IDs are text like '001'. Reading it must not
# turn '001' into 1.
def two_sheets() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "kunder"
    ws.append(["kunde_id", "navn", "saldo"])
    for row in [("001", "Alpha", 1000), ("002", "Beta", 2000), ("003", "Gamma", 3000)]:
        ws.append(row)
    for cell in ws["A"][1:]:
        cell.number_format = "@"

    ws2 = wb.create_sheet("konti")
    ws2.append(["konto_id", "kunde_id", "type"])
    for row in [("A1", "001", "loen"), ("A2", "002", "opsparing")]:
        ws2.append(row)

    wb.save(HERE / "two_sheets.xlsx")
    print("wrote two_sheets.xlsx")


# 10 — a Parquet file that carries its own types. The loader must keep them
# rather than stringify everything, which is the documented exception to
# loading as literal text (spec 4.2).
def typed_parquet() -> None:
    import duckdb

    path = (HERE / "typed.parquet").as_posix()
    con = duckdb.connect()
    con.execute(
        """
        COPY (
            SELECT * FROM (VALUES
                ('001', 31, 12345.50, DATE '2021-04-03', TRUE),
                ('002', 54,  7800.00, DATE '2020-02-28', FALSE),
                ('003', 23, 250000.0, DATE '2022-11-13', TRUE),
                ('004', NULL, NULL,   DATE '2019-01-30', FALSE)
            ) AS t(kunde_id, alder, beloeb, dato, aktiv)
        ) TO ? (FORMAT PARQUET)
        """,
        [path],
    )
    con.close()
    print("wrote typed.parquet")


GENERATORS = [
    danish_semicolon,
    bom_excel_export,
    latin1_aeoeaa,
    quoted_newline,
    duplicate_blank_headers,
    too_many_fields,
    ambiguous_dates,
    cpr_numbers,
    two_sheets,
    typed_parquet,
]


def main() -> None:
    for generator in GENERATORS:
        generator()


if __name__ == "__main__":
    main()
